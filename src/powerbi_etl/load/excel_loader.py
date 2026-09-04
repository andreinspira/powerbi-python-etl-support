"""
Excel/CSV loader for Power BI data sources.
Outputs formatted Excel files ready for Power BI consumption.
"""

import pandas as pd
from pathlib import Path
from typing import Dict, List, Optional, Any
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils.dataframe import dataframe_to_rows
from .base import BaseLoader, LoadResult
import structlog

logger = structlog.get_logger(__name__)


class ExcelLoader(BaseLoader):
    """
    Loads dataframes to formatted Excel files.
    Produces files optimized for Power BI consumption with proper formatting.
    """

    def __init__(self, config: Dict[str, Any]):
        super().__init__(config)
        self.output_path = Path(config.get("output_path", "data/output/etl_output.xlsx"))
        self.sheet_name = config.get("sheet_name", "Data")
        self.include_index = config.get("include_index", False)
        self.freeze_panes = config.get("freeze_panes", "A2")
        self.auto_filter = config.get("auto_filter", True)
        self.column_formats = config.get("column_formats", {})
        self.add_metadata_sheet = config.get("add_metadata_sheet", True)

    def test_connection(self) -> bool:
        """Test if output directory is writable."""
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            test_file = self.output_path.parent / ".write_test"
            test_file.write_text("test")
            test_file.unlink()
            return True
        except Exception as e:
            self.logger.error("Connection test failed", error=str(e))
            return False

    def load(self, dataframe: "pd.DataFrame") -> LoadResult:
        """Load dataframe to formatted Excel file."""
        try:
            self.logger.info("Loading to Excel", path=str(self.output_path), rows=len(dataframe))
            
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            
            wb = Workbook()
            ws = wb.active
            ws.title = self.sheet_name
            
            # Write data
            for r_idx, row in enumerate(dataframe_to_rows(dataframe, index=self.include_index, header=True), 1):
                for c_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=value)
                    
                    # Header formatting
                    if r_idx == 1:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", wrap_text=True)
            
            # Apply column formats
            for col_idx, col_name in enumerate(dataframe.columns, 1):
                if col_name in self.column_formats:
                    fmt = self.column_formats[col_name]
                    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                        for cell in row:
                            if fmt.get("number_format"):
                                cell.number_format = fmt["number_format"]
                            if fmt.get("alignment"):
                                cell.alignment = Alignment(**fmt["alignment"])
            
            # Freeze panes
            if self.freeze_panes:
                ws.freeze_panes = self.freeze_panes
            
            # Auto filter
            if self.auto_filter and len(dataframe) > 0:
                ws.auto_filter.ref = ws.dimensions
            
            # Adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Add metadata sheet
            if self.add_metadata_sheet:
                self._add_metadata_sheet(wb, dataframe)
            
            wb.save(self.output_path)
            
            self.logger.info("Excel file saved", path=str(self.output_path), rows=len(dataframe))
            
            return self._create_result(
                success=True,
                records_loaded=len(dataframe),
                metadata={
                    "file_path": str(self.output_path),
                    "sheet_name": self.sheet_name,
                    "rows": len(dataframe),
                    "columns": len(dataframe.columns),
                },
            )
            
        except Exception as e:
            self.logger.error("Excel load failed", error=str(e))
            return self._create_result(success=False, errors=[str(e)])

    def _add_metadata_sheet(self, wb: Workbook, dataframe: "pd.DataFrame") -> None:
        """Add metadata/documentation sheet."""
        meta_ws = wb.create_sheet("ETL_Metadata")
        
        meta_data = [
            ["Property", "Value"],
            ["Generated At", datetime.now().isoformat()],
            ["Source Rows", len(dataframe)],
            ["Source Columns", len(dataframe.columns)],
            ["Column Names", ", ".join(dataframe.columns)],
            ["Data Types", ", ".join([f"{c}: {str(dataframe[c].dtype)}" for c in dataframe.columns])],
            ["Memory Usage (MB)", round(dataframe.memory_usage(deep=True).sum() / 1024 / 1024, 2)],
            ["Null Counts", ", ".join([f"{c}: {int(dataframe[c].isna().sum())}" for c in dataframe.columns])],
        ]
        
        for r_idx, row in enumerate(meta_data, 1):
            for c_idx, value in enumerate(row, 1):
                cell = meta_ws.cell(row=r_idx, column=c_idx, value=value)
                if r_idx == 1:
                    cell.font = Font(bold=True)
        
        # Auto-fit
        for column in meta_ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            meta_ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 80)


class MultiSheetExcelLoader(BaseLoader):
    """Load multiple dataframes to separate sheets in one Excel file."""

    def __init__(self, config: Dict[str, Any]):
        super().__init__(config)
        self.output_path = Path(config.get("output_path", "data/output/etl_multi_sheet.xlsx"))
        self.sheets_config = config.get("sheets", {})  # {sheet_name: dataframe_key}

    def test_connection(self) -> bool:
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            return True
        except Exception:
            return False

    def load(self, dataframes: Dict[str, "pd.DataFrame"]) -> LoadResult:
        """Load multiple dataframes to separate sheets."""
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            
            with pd.ExcelWriter(self.output_path, engine="openpyxl") as writer:
                for sheet_name, df in dataframes.items():
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    ws = writer.sheets[sheet_name]
                    self._format_sheet(ws, df)
            
            total_rows = sum(len(df) for df in dataframes.values())
            self.logger.info("Multi-sheet Excel saved", path=str(self.output_path), sheets=len(dataframes), total_rows=total_rows)
            
            return self._create_result(
                success=True,
                records_loaded=total_rows,
                metadata={
                    "file_path": str(self.output_path),
                    "sheets": list(dataframes.keys()),
                    "total_rows": total_rows,
                },
            )
        except Exception as e:
            self.logger.error("Multi-sheet Excel load failed", error=str(e))
            return self._create_result(success=False, errors=[str(e)])

    def _format_sheet(self, ws, df: "pd.DataFrame") -> None:
        # Header formatting
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        ws.freeze_panes = "A2"
        if len(df) > 0:
            ws.auto_filter.ref = ws.dimensions


class CSVLoader(BaseLoader):
    """Load dataframes to CSV files."""

    def __init__(self, config: Dict[str, Any]):
        super().__init__(config)
        self.output_path = Path(config.get("output_path", "data/output/etl_output.csv"))
        self.encoding = config.get("encoding", "utf-8")
        self.delimiter = config.get("delimiter", ",")
        self.include_index = config.get("include_index", False)
        self.date_format = config.get("date_format", "%Y-%m-%d")

    def test_connection(self) -> bool:
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            return True
        except Exception:
            return False

    def load(self, dataframe: "pd.DataFrame") -> LoadResult:
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Format dates
            df = dataframe.copy()
            for col in df.select_dtypes(include=["datetime64", "datetimetz"]).columns:
                df[col] = df[col].dt.strftime(self.date_format)
            
            df.to_csv(
                self.output_path,
                index=self.include_index,
                encoding=self.encoding,
                sep=self.delimiter,
            )
            
            self.logger.info("CSV file saved", path=str(self.output_path), rows=len(df))
            
            return self._create_result(
                success=True,
                records_loaded=len(df),
                metadata={"file_path": str(self.output_path), "rows": len(df)},
            )
        except Exception as e:
            self.logger.error("CSV load failed", error=str(e))
            return self._create_result(success=False, errors=[str(e)])


class ParquetLoader(BaseLoader):
    """Load dataframes to Parquet files (efficient for intermediate storage)."""

    def __init__(self, config: Dict[str, Any]):
        super().__init__(config)
        self.output_path = Path(config.get("output_path", "data/processed/etl_output.parquet"))
        self.compression = config.get("compression", "snappy")

    def test_connection(self) -> bool:
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            return True
        except Exception:
            return False

    def load(self, dataframe: "pd.DataFrame") -> LoadResult:
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            dataframe.to_parquet(self.output_path, compression=self.compression, index=False)
            
            self.logger.info("Parquet file saved", path=str(self.output_path), rows=len(dataframe))
            
            return self._create_result(
                success=True,
                records_loaded=len(dataframe),
                metadata={"file_path": str(self.output_path), "rows": len(dataframe)},
            )
        except Exception as e:
            self.logger.error("Parquet load failed", error=str(e))
            return self._create_result(success=False, errors=[str(e)])